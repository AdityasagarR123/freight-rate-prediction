"""Build report/Freight_Rate_Report.docx from the outputs of validate.py, train.py and score.py.

Usage: python src/make_report.py
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"

MODEL_LABELS = {
    "hybrid_gbm_time_terms": "Hybrid: LightGBM (load features) + linear trend & quarter-end terms  [chosen]",
    "hybrid_gbm_time_terms_no_outlier_filter": "Hybrid, trained without label-outlier filter",
    "lightgbm_detrended": "LightGBM with calendar features, linear detrend",
    "lightgbm": "LightGBM with calendar features, no trend handling",
    "huber_log_linear": "Robust (Huber) log-linear baseline",
}


def table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    grid = doc.add_table(rows=1, cols=len(header))
    grid.style = "Light Grid Accent 1"
    for cell, text in zip(grid.rows[0].cells, header):
        cell.text = text
    for row in rows:
        for cell, text in zip(grid.add_row().cells, row):
            cell.text = text
    for row in grid.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def main() -> None:
    cv = pd.read_csv(OUT / "cv_results.csv")
    meta = json.loads((OUT / "model_meta.json").read_text())
    importance = pd.read_csv(OUT / "feature_importance.csv", index_col=0)["gain"]
    summary = cv.drop(columns="fold").groupby("model").mean().sort_values("clean_mae")
    chosen = summary.loc["hybrid_gbm_time_terms"]
    baseline = summary.loc["huber_log_linear"]

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading("Freight Rate Prediction: Approach & Validation", 0)
    doc.add_paragraph(
        "Goal: predict posted_rate for 12,000 loads dated Nov-Dec 2025 using 48,000 labelled loads "
        "dated Jan-Oct 2025. The final model is a LightGBM regressor on log(rate) using load-level "
        "features, combined with two linear time terms (a long-run trend and a quarter-end ramp) "
        "that extrapolate into the unseen months."
    )
    doc.add_paragraph(
        f"Out-of-time backtest (mean of three 2-month folds): MAPE {chosen.clean_mape:.2f}% / "
        f"MAE ${chosen.clean_mae:.0f} on clean loads, versus {baseline.clean_mape:.2f}% / "
        f"${baseline.clean_mae:.0f} for a robust log-linear baseline. Including corrupted labels, "
        f"MAPE is {chosen.all_mape:.2f}% and median APE {chosen.all_mdape:.2f}%."
    )

    doc.add_heading("1. Key findings from the data", 1)
    bullets(doc, [
        "The split is by time: train_test covers 2025-01-01 to 2025-10-31; validation covers "
        "2025-11-01 to 2025-12-31. Validation is a forecast of the next two months, so random "
        "K-fold splits would leak future market conditions and overstate accuracy.",
        "Distance explains most of the variance (rate per mile falls with distance). Reefer costs "
        "~12% and Flatbed ~8% more than Dry Van; heavier loads cost slightly more (~3% per 10,000 lb).",
        "market_index is almost a per-day series: its spread within a day (std 0.025) is far "
        "smaller than its spread between days (0.17). A 10% rise in the index adds ~1.5% to the rate.",
        "Rates ramp up roughly linearly over the last ~3 weeks of every quarter (about +3.5% by "
        "quarter end in March, June and September) and drop back on the first day of the next "
        "quarter. December 31 is a quarter end, so this matters for the December forecast.",
        "There is an upward drift of about 7% per year that market_index does not explain.",
        "quote_signal is a weak signal. It adds a little through non-linear effects.",
        "Coordinates are synthetic/jittered (e.g. New York to Allentown is 0.9 miles apart), and "
        "distance has a 70-mile floor, so posted distance, not haversine, is the authoritative "
        "length. Rates line up with posted distance even where it disagrees with the coordinates.",
    ])

    doc.add_heading("2. Data-quality issues and how they were handled", 1)
    table(doc, ["Issue", "Extent", "Handling"], [
        ["Corrupted labels: posted_rate at ~0.2-0.5x or ~2-6x the expected rate",
         f"~1.4% of rows ({meta['label_outliers_removed']} flagged)",
         "Fit a robust Huber log-linear model and drop rows with |log residual| > 0.25 "
         "(clean residual std is ~0.04, so this cut is above 6 sigma). The filter is fitted "
         "only on each training fold."],
        ["Negative weights", "292 train / 145 validation rows",
         "Treated as sign errors and replaced with |weight|, which falls in the normal range."],
        ["Missing weight", "300 train / 165 validation", "Equipment median, plus a missing-value flag."],
        ["Missing market_index", "374 train / 249 validation",
         "Filled with the same day's mean index, since the index is effectively a daily series."],
        ["Cities unseen in training", "8 cities (Chicago, Charlotte, Laredo, ...) in validation",
         "City identities are not used as features. Location enters through lat/lon, which "
         "generalise to new cities."],
        ["Distance vs coordinates mismatch", "~200 rows with ratio > 1.5 (up to 78x)",
         "Investigated and kept: rates follow posted distance, and the mismatches come from the "
         "70-mile floor and jittered coordinates."],
    ])

    doc.add_heading("3. Train/test split and validation approach", 1)
    doc.add_paragraph(
        "The validation set is the two months after the training data ends, so the backtest mimics "
        "that exactly. It uses expanding-window, out-of-time folds: each fold trains on all data "
        "before a cutoff and tests on the following two months. Folds F1 and F3 both contain a "
        "quarter end, which tests whether the quarter-end ramp carries over."
    )
    table(doc, ["Fold", "Train window", "Test window"], [
        ["F1", "Jan-Apr 2025", "May-Jun 2025"],
        ["F2", "Jan-Jun 2025", "Jul-Aug 2025"],
        ["F3", "Jan-Aug 2025", "Sep-Oct 2025"],
    ])
    doc.add_paragraph(
        "Metrics are reported on all test loads and on 'clean' test loads. Clean loads exclude "
        "corrupted labels, flagged by the filter fitted on the training fold. No model can "
        "predict a 5x label error, so the clean metrics compare model quality fairly, while the "
        "all-rows metrics show what to expect against the raw labels. All preprocessing is "
        "fitted inside each fold. Daily market means use only same-day feature values, never labels."
    )
    doc.add_paragraph("Backtest results (mean over F1-F3):")
    rows = []
    for name, row in summary.iterrows():
        rows.append([MODEL_LABELS.get(name, name), f"{row.clean_mae:.1f}", f"{row.clean_mape:.2f}%",
                     f"{row.all_mae:.1f}", f"{row.all_mape:.2f}%", f"{row.all_mdape:.2f}%"])
    table(doc, ["Model", "Clean MAE $", "Clean MAPE", "All MAE $", "All MAPE", "All MdAPE"], rows)
    doc.add_paragraph("")
    per_fold = cv[cv.model == "hybrid_gbm_time_terms"]
    table(doc, ["Chosen model by fold", "Clean MAE $", "Clean MAPE", "All MAPE"],
          [[r.fold, f"{r.clean_mae:.1f}", f"{r.clean_mape:.2f}%", f"{r.all_mape:.2f}%"] for r in per_fold.itertuples()])
    doc.add_paragraph(
        "F2 (Jul-Aug) is the hardest fold. Rates dipped about 1% below what the market index "
        "and trend imply, a market-wide move that no feature in the data anticipates."
    )
    doc.add_picture(str(OUT / "holdout_daily.png"), width=Inches(6.3))

    doc.add_heading("4. Model choice", 1)
    bullets(doc, [
        "The target is log(rate). Errors are multiplicative, and exponentiating a log prediction "
        "gives a median-style estimate that is robust for MAE/MAPE.",
        "A plain GBM given calendar features (day of quarter, daily market mean) memorises "
        "individual training days. It cannot extrapolate time, so its F3 error was nearly double.",
        "The hybrid keeps the GBM for what it does well: non-linear distance curves, "
        "equipment x distance interactions and regional (lat/lon) effects. Calendar effects that "
        "must carry into Nov-Dec (trend, quarter-end ramp) are linear terms. The two parts are "
        "fitted by backfitting with a robust Huber loss.",
        f"Final time terms: trend {meta['time_coefficients']['trend_per_year']*100:.1f}% per year, "
        f"quarter-end ramp +{meta['time_coefficients']['quarter_end_ramp']*100:.1f}% at quarter end. "
        "LightGBM uses a Huber objective, learning rate 0.03, 2,500 trees and 31 leaves.",
        "Removing corrupted labels before training lowered error slightly in every fold.",
        "Feature importance (gain share): " + ", ".join(f"{k} {v*100:.1f}%" for k, v in importance.head(6).items()) + ".",
    ])

    doc.add_heading("5. December 2025 fixed-lane chart", 1)
    doc.add_paragraph(
        "The chart is produced by the provided score.py. The December input rows have no "
        "market_index or quote_signal, so each date uses the daily mean of the validation loads "
        "on that date (the only market information for those days). The curve shows the "
        "market index's weekly cycle and the quarter-end ramp lifting rates through late December."
    )
    doc.add_picture(str(ROOT / "scorer_results" / "candidate_december.png"), width=Inches(6.5))

    doc.add_heading("6. Limitations and next steps", 1)
    bullets(doc, [
        "Two-month-ahead market moves not reflected in market_index (as in Jul-Aug) remain the "
        "largest error source. A leading market indicator would help most.",
        "Nov-Dec holidays (Thanksgiving, Christmas) do not appear in the training months, so "
        "their effect cannot be learned from this data.",
        "The trend is linear and estimated from ten months. It should be re-estimated as new "
        "data arrives.",
    ])

    target = ROOT / "report" / "Freight_Rate_Report.docx"
    target.parent.mkdir(exist_ok=True)
    doc.save(target)
    print(f"Wrote {target}")


if __name__ == "__main__":
    main()
